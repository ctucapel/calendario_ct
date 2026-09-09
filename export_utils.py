import os
import re
import shutil
import sqlite3
import subprocess
import tempfile
import unicodedata
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

DUOC_YELLOW = '#F1B634'
DUOC_BLACK = '#1A1A1A'
DUOC_GRAY = '#666666'

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_2027 = os.path.join(BASE_DIR, 'templates', 'Calendario_Academico_Base_2027.xlsx')


def _df(con, year, version_id=None):
    if version_id is not None:
        return pd.read_sql_query('''
          SELECT a.id AS ID, a.code AS Código, g.name AS Grupo, a.name AS Actividad,
                 vo.semester AS Semestre, vo.start_date AS Inicio, vo.end_date AS Término,
                 a.observations AS Observaciones
          FROM version_occurrences vo
          JOIN activities a ON a.id=vo.activity_id
          JOIN activity_groups g ON g.id=a.group_id
          WHERE vo.version_id=? AND (vo.semester LIKE ? OR vo.semester='TAV')
          ORDER BY g.id,a.id,vo.semester''', con, params=(version_id, f'{year}%'))
    return pd.read_sql_query('''
      SELECT a.id AS ID, a.code AS Código, g.name AS Grupo, a.name AS Actividad, o.semester AS Semestre,
             o.start_date AS Inicio, o.end_date AS Término, a.observations AS Observaciones
      FROM occurrences o JOIN activities a ON a.id=o.activity_id JOIN activity_groups g ON g.id=a.group_id
      WHERE o.semester LIKE ? OR o.semester='TAV'
      ORDER BY g.id,a.id,o.semester''', con, params=(f'{year}%',))


def _norm(text):
    text = '' if text is None else str(text)
    text = unicodedata.normalize('NFKD', text)
    text = ''.join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r'\s+', ' ', text).strip().lower()


def _excel_date(value):
    if value in (None, '', '-'):
        return None
    if isinstance(value, datetime):
        return value
    return datetime.strptime(str(value)[:10], '%Y-%m-%d')


def _build_occurrence_map(df):
    out = {}
    for _, row in df.iterrows():
        out[(_norm(row['Actividad']), str(row['Semestre']))] = row
    return out


def _update_process_sheet(wb, df, year):
    """Actualiza fechas en CA Procesos sin alterar formato, bordes, colores ni tamaños."""
    ws = wb['CA Procesos']
    occ = _build_occurrence_map(df)
    semester_cols = {
        'TAV': ('D', 'E'),
        f'{year}-1': ('F', 'G'),
        f'{year}-2': ('H', 'I'),
    }
    for row_num in range(8, ws.max_row + 1):
        activity_name = ws[f'C{row_num}'].value
        key_name = _norm(activity_name)
        if not key_name:
            continue
        for sem, (start_col, end_col) in semester_cols.items():
            item = occ.get((key_name, sem))
            if item is None:
                continue
            start = _excel_date(item['Inicio'])
            end = _excel_date(item['Término'])
            # Si la ocurrencia existe, sus valores mandan. Para fecha ausente usamos celda vacía.
            ws[f'{start_col}{row_num}'] = start
            # Algunas actividades de un solo día usan celdas combinadas (Inicio-Término).
            end_cell = ws[f'{end_col}{row_num}']
            if end_cell.__class__.__name__ != 'MergedCell':
                end_cell.value = end

    # Mantener el título coherente para futuras plantillas derivadas.
    if isinstance(ws['C2'].value, str):
        ws['C2'] = re.sub(r'CALENDARIO ACADÉMICO\s+\d{4}', f'CALENDARIO ACADÉMICO {year}', ws['C2'].value)


def _configure_printing(wb):
    """Tres formatos, tres páginas: Procesos, Cronológico y Gráfico."""
    if 'CA Gráfico 2025' in wb.sheetnames:
        del wb['CA Gráfico 2025']

    configs = {
        'CA Procesos': ('portrait', 'C2:I69'),
        'CA Cronológico': ('portrait', 'D2:F93'),
        'CA Gráfico 2027': ('landscape', 'D5:AJ42'),
    }
    for name, (orientation, area) in configs.items():
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        ws.sheet_state = 'visible'
        ws.print_area = area
        ws.page_setup.orientation = orientation
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = 0.15
        ws.page_margins.right = 0.15
        ws.page_margins.top = 0.20
        ws.page_margins.bottom = 0.20
        ws.page_margins.header = 0.10
        ws.page_margins.footer = 0.10
    wb.active = wb.sheetnames.index('CA Procesos') if 'CA Procesos' in wb.sheetnames else 0


def _make_target_workbook(source_path, out_path, target_sheet):
    """Conserva hojas de soporte para fórmulas, pero muestra solo el formato descargado."""
    wb = load_workbook(source_path)
    if 'CA Gráfico 2025' in wb.sheetnames:
        del wb['CA Gráfico 2025']
    for ws in wb.worksheets:
        ws.sheet_state = 'visible' if ws.title == target_sheet else 'hidden'
    wb.active = wb.sheetnames.index(target_sheet)
    wb.save(out_path)


def _libreoffice_pdf(xlsx_path, pdf_path):
    exe = shutil.which('libreoffice') or shutil.which('soffice')
    if not exe:
        raise RuntimeError('LibreOffice no está instalado. En Streamlit Cloud debe existir packages.txt con libreoffice-calc.')
    out_dir = os.path.dirname(pdf_path)
    os.makedirs(out_dir, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='lo_profile_') as profile:
        profile_uri = 'file://' + profile
        cmd = [
            exe, '--headless', '--norestore', '--nofirststartwizard',
            f'-env:UserInstallation={profile_uri}',
            '--convert-to', 'pdf', '--outdir', out_dir, xlsx_path,
        ]
        result = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=90)
        generated = os.path.join(out_dir, os.path.splitext(os.path.basename(xlsx_path))[0] + '.pdf')
        if result.returncode != 0 or not os.path.exists(generated):
            raise RuntimeError(f'No fue posible generar PDF con LibreOffice. {result.stdout[-1200:]}')
        if os.path.abspath(generated) != os.path.abspath(pdf_path):
            if os.path.exists(pdf_path):
                os.remove(pdf_path)
            os.replace(generated, pdf_path)


def _prepare_exact_workbook(df, year, out_path):
    if int(year) != 2027 or not os.path.exists(TEMPLATE_2027):
        raise RuntimeError('El formato institucional exacto está configurado actualmente para el Calendario 2027.')
    wb = load_workbook(TEMPLATE_2027)
    _update_process_sheet(wb, df, year)
    _configure_printing(wb)
    # Fuerza recálculo de fórmulas al abrir con Excel/LibreOffice.
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = 'auto'
    except Exception:
        pass
    wb.save(out_path)


def export_all(db_path, year, version_no, out_dir='exports', version_id=None):
    """
    Genera los tres formatos originales de Excel y un PDF de tres páginas.
    Página 1: CA Procesos (vertical)
    Página 2: CA Cronológico (vertical)
    Página 3: CA Gráfico 2027 (horizontal)
    """
    os.makedirs(out_dir, exist_ok=True)
    con = sqlite3.connect(db_path)
    df = _df(con, year, version_id)
    con.close()

    with tempfile.TemporaryDirectory(prefix='cal_export_') as tmp:
        master = os.path.join(tmp, f'Calendario_{year}_V{version_no}_master.xlsx')
        _prepare_exact_workbook(df, int(year), master)

        p1 = os.path.join(out_dir, f'CA_Procesos_{year}_V{version_no}.xlsx')
        p2 = os.path.join(out_dir, f'CA_Cronologico_{year}_V{version_no}.xlsx')
        p3 = os.path.join(out_dir, f'CA_Grafico_{year}_V{version_no}.xlsx')
        _make_target_workbook(master, p1, 'CA Procesos')
        _make_target_workbook(master, p2, 'CA Cronológico')
        _make_target_workbook(master, p3, 'CA Gráfico 2027')

        # El PDF se genera desde un libro que contiene exactamente las tres hojas visibles.
        pdf_book = os.path.join(tmp, f'Calendario_Academico_{year}_V{version_no}.xlsx')
        shutil.copy2(master, pdf_book)
        pdf = os.path.join(out_dir, f'CA_Calendario_{year}_V{version_no}.pdf')
        _libreoffice_pdf(pdf_book, pdf)

    return [p1, p2, p3, pdf]
